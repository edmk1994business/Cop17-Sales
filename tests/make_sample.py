"""Write a mock COP-17-Sales.xlsx that mimics the iiko OLAP pivot export.

Usage:  python tests/make_sample.py [out.xlsx]

Each sheet has 4 title rows, a header row, merged-style blanks (the date is written
only on the first row of its block), a "dd.mm.yyyy всего" row after every date and
an "Итого" grand total — exactly the shape the parser has to survive.
Sheet variants exercise different real-world quirks:
  LOT 3  – standard layout, "Чеков", string dates, discount column with zeros
  LOT 8  – an extra outer "Торговое предприятие" column + its own "… всего" rows
  LOT 9  – datetime cells, "Заказов" instead of "Чеков", numbers as strings "12 345,00"
  LOT 10 – real discounts, plus "(без оплаты)" compliments rows
Column-C summary sheets (one total in column C + the period written on the sheet):
  LOT 4                          – one row: label · "19.10.2026 - 30.10.2026" · 5 400 000
  side events, CBD, Ֆուրշետներ    – dated entry rows + an "Ընդամենը" total row
  պետական պատվերներ               – period as "19-30.10.2026", total as text "3 200 000", a "Чеков" row
"""
from __future__ import annotations

import sys
from datetime import date, datetime, timedelta

import numpy as np
from openpyxl import Workbook

START = date(2026, 10, 19)
DAYS = 4  # 19–22 Oct
PAYMENTS = ["POS Arcus", "Банковские карты", "Наличные"]


def rows_for(seed: int, with_discount: bool, compliments: bool):
    rng = np.random.default_rng(seed)
    out = []
    for i in range(DAYS):
        d = START + timedelta(days=i)
        lines = []
        for p in PAYMENTS:
            checks = int(rng.integers(20, 90))
            sales = int(rng.integers(3000, 6000)) * checks
            disc = int(sales * 0.07) if with_discount else 0
            lines.append((p, checks, sales, disc))
        if compliments:
            lines.append(("(без оплаты)", int(rng.integers(1, 4)), int(rng.integers(8000, 30000)), 0))
        out.append((d, lines))
    return out


def write_standard(ws, title, data, checks_label="Чеков", as_datetime=False, as_text=False,
                   entity=None):
    ws.append([title]); ws.append(["Период: 19.10.2026 - 30.10.2026"]); ws.append([]); ws.append([])
    header = ["Учетный день", "Тип оплаты", checks_label, "Сумма со скидкой", "Сумма скидки"]
    if entity:
        header = ["Торговое предприятие"] + header
    ws.append(header)
    tot_c = tot_s = tot_d = 0
    for d, lines in data:
        day_c = day_s = day_d = 0
        for k, (p, c, s, dsc) in enumerate(lines):
            if as_datetime:
                dcell = datetime(d.year, d.month, d.day) if k == 0 else None
            else:
                dcell = d.strftime("%d.%m.%Y") if k == 0 else None
            vals = [c, s, dsc]
            if as_text:
                vals = [str(c), f"{s:,.2f}".replace(",", " ").replace(".", ","), str(dsc)]
            row = [dcell, p, *vals]
            if entity:
                row = [entity if (k == 0 and d == data[0][0]) else None] + row
            ws.append(row)
            day_c += c; day_s += s; day_d += dsc
        sub = [f"{d.strftime('%d.%m.%Y')} всего", None, day_c, day_s, day_d]
        if entity:
            sub = [None] + sub
        ws.append(sub)
        tot_c += tot_c * 0 + day_c; tot_s += day_s; tot_d += day_d
    if entity:
        ws.append([f"{entity} всего", None, None, tot_c, tot_s, tot_d])
    tot = ["Итого", None, tot_c, tot_s, tot_d]
    if entity:
        tot = [None] + tot
    ws.append(tot)
    return tot_c, tot_s, tot_d


def write_period_sheets(wb):
    ws = wb.create_sheet("LOT 4")
    ws.append(["LOT 4 — COP17 catering"]); ws.append([])
    ws.append(["Период", "19.10.2026 - 30.10.2026", 5_400_000])

    ws = wb.create_sheet("side events, CBD, Ֆուրշետներ")
    ws.append(["Ամսաթիվ", "Միջոցառում", "Գումար"])
    ws.append([datetime(2026, 10, 20), "CBD reception", 850_000])
    ws.append([datetime(2026, 10, 22), "Ֆուրշետ · Pavilion B", 1_250_000])
    ws.append([datetime(2026, 10, 24), "Side event dinner", 640_000])
    ws.append(["Ընդամենը", None, 2_740_000])

    ws = wb.create_sheet("պետական պատվերներ")
    ws.append(["Պետական պատվերներ"]); ws.append(["Ժամանակահատված", "19-30.10.2026", "3 200 000"])
    ws.append(["Чеков", None, 14])
    return {"LOT 4": (None, 5_400_000, (date(2026, 10, 19), date(2026, 10, 30))),
            "side events, CBD, Ֆուրշետներ": (None, 2_740_000, (date(2026, 10, 20), date(2026, 10, 24))),
            "պետական պատվերներ": (14, 3_200_000, (date(2026, 10, 19), date(2026, 10, 30)))}


def main(path: str):
    wb = Workbook()
    wb.remove(wb.active)
    expected = {}
    expected["LOT 3"] = write_standard(wb.create_sheet("LOT 3"), "OLAP Отчет по продажам — LOT 3",
                                       rows_for(3, False, False))
    expected["LOT 8"] = write_standard(wb.create_sheet("LOT 8"), "OLAP Отчет по продажам — LOT 8",
                                       rows_for(8, False, False), entity="COP17 LOT 8")
    expected["LOT 9"] = write_standard(wb.create_sheet("LOT 9"), "OLAP Отчет по продажам — LOT 9",
                                       rows_for(9, False, False), checks_label="Заказов",
                                       as_datetime=True, as_text=True)
    expected["LOT 10"] = write_standard(wb.create_sheet("LOT 10"), "OLAP Отчет по продажам — LOT 10",
                                        rows_for(10, True, True))
    periods = write_period_sheets(wb)
    wb.save(path)
    print("wrote", path)
    for k, (c, s, d) in expected.items():
        print(f"  {k:7} checks={c:5} sales={s:>12,} discount={d:>10,}")
    for k, (c, s, (a, b)) in periods.items():
        print(f"  {k:30} total={s:>12,} checks={c} period={a}..{b}")
    return expected, periods


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "data/COP-17-Sales.xlsx")
